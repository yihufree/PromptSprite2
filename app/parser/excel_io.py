# -*- coding: utf-8 -*-
"""
excel_io.py - Excel(.xlsx) 数据导入导出（批量编辑场景）
创建日期：2026-08-12（阶段六创建；阶段八重构为全局分类+领域关联 v2）

导出列：领域、一级分类、二级分类、名称、介绍、溯源、核心特征、应用场景、
        代表作、配图描述、提示词中文、提示词英文、图像获取方案、收藏
导入：按 领域/一级/二级 名称解析或新建全局分类，并建立 领域↔一级分类 关联（多对一共享）。
"""
from openpyxl import Workbook, load_workbook

from .. import config  # 2026-08-29（B3 修复）：新建根目录兜底归入"未明确分类"
from ..database import Database  # 2026-08-18（P1-1）：内容判重键 content_key
from ..models import Entry
from .json_io import (_gather_categories, _chain_names, _restore_pending_custom_fields,
                      _restore_pending_tags,
                      _restore_pending_images)

_HEADERS = ["领域", "一级分类", "二级分类", "名称", "介绍", "溯源", "核心特征",
            "应用场景", "代表作", "配图描述", "提示词中文", "提示词英文", "图像获取方案", "收藏"]

# 2026-09-13（1-A-5 收尾 / 1-C-4）：追加列的表头
#   · 自定义字段列："自定义：<显示名>"（导入按显示名匹配既有字段）
#   · 标签列："标签"（多个用"、"分隔；导入会自动创建缺失标签）
_CUSTOM_PREFIX = "自定义："
_TAG_HEADER = "标签"
# 2026-09-13（2-d）：图集列（本地图记相对路径、外链图记网址，多个用"、"分隔）
_GALLERY_HEADER = "图集"
# 2026-09-13（追加列收尾）：项目类别列（Excel 原只有 领域/一级/二级，缺最高层；
#   本列让 Excel 往返也能恢复"项目类别"归属；旧表无此列时按原逻辑处理，完全兼容）
_PROJECT_HEADER = "项目类别"


def _entry_project_name(db, entry) -> str:
    """条目所属根目录的项目类别名（取第一个关联根目录；未分配/异常返回空串）"""
    try:
        root = db.category_root(entry["category_id"]) if entry["category_id"] else None
        if root is None:
            return ""
        for d in db.linked_domains(root):
            pid = d.get("project_id")
            if pid:
                p = db.get_project(pid)
                if p:
                    return p["name"]
    except Exception:
        pass
    return ""


def _gallery_text(db, entry_id: int) -> str:
    """条目图集 → 导出文本（本地图=相对路径，外链图=网址；"、"分隔）"""
    try:
        items = db.list_entry_images(entry_id)
    except Exception:
        items = []
    out = []
    for g in items:
        if g.get("kind") == "url":
            if g.get("source_url"):
                out.append(str(g["source_url"]))
        elif g.get("path"):
            out.append(str(g["path"]))
    return "、".join(out)


def _split_gallery(text: str) -> list:
    """图集文本 → 图集项列表（http(s) 视为外链，其余视为本地相对路径）"""
    if not text:
        return []
    out = []
    for raw in str(text).replace("|", "、").replace("；", "、").replace(";", "、") \
            .replace("，", "、").replace(",", "、").split("、"):
        t = raw.strip()
        if not t:
            continue
        if t.startswith(("http://", "https://")):
            out.append({"kind": "url", "path": "", "source_url": t, "caption": ""})
        else:
            out.append({"kind": "local", "path": t, "source_url": "", "caption": ""})
    return out


def _tags_text(db, entry_id: int) -> str:
    """条目标签 → 导出文本（"、"分隔）"""
    try:
        return "、".join(db.list_entry_tag_names(entry_id))
    except Exception:
        return ""


def _split_tags(text: str) -> list:
    """标签文本 → 名称列表（兼容 "、,;；|" 分隔与空白）"""
    if not text:
        return []
    out = []
    for raw in str(text).replace("|", "、").replace("；", "、").replace(";", "、") \
            .replace("，", "、").replace(",", "、").split("、"):
        n = raw.strip()
        if n and n not in out:
            out.append(n)
    return out


def _custom_values(db, entry_id: int, custom_defs: list) -> list:
    """取某条目在给定自定义字段定义下的取值列表（顺序与 custom_defs 一致）"""
    try:
        vals = {r["field_key"]: r["value_text"]
                for r in db.list_entry_field_values(entry_id)}
    except Exception:
        vals = {}
    return [vals.get(d["field_key"], "") or "" for d in custom_defs]


def resolve_category_path(db, domain_name: str, l1_name: str, l2_name: str,
                          project_name: str = ""):
    """按 项目类别/领域/一级/二级 名称解析目标分类 id（多对一共享）：

    - 领域不存在则创建；一级分类已存在则仅建立领域关联，不存在则新建并关联；
    - 二级分类不存在则在一级下新建；全空返回 None(未分类)；
    - project_name（2026-09-13，4-c 追加列）：非空时把该领域归入指定项目类别（不存在则新建），
      为空时新领域沿用既有兜底（未明确分类），已有领域不改归属。
    """
    if not domain_name:
        return None
    domain = next((d for d in db.list_domains() if d["name"] == domain_name), None)
    if domain:
        did = domain["id"]
        if project_name:
            db.move_domain_to_project(did, db.ensure_project(project_name))
    else:
        # 2026-08-29（B3 修复）：新建根目录默认归入"未明确分类"，避免落"未分配"
        fallback = (db.ensure_project(project_name) if project_name
                    else db.ensure_project(config.PROJECT_FALLBACK))
        did = db.add_domain(domain_name, project_id=fallback)
    if not l1_name:
        return None
    l1 = next((c for c in db.list_categories(parent_id=None) if c["name"] == l1_name), None)
    if l1 is None:
        l1_id = db.add_category(l1_name, domain_id=did)
    else:
        l1_id = l1["id"]
        db.link_domain_category(did, l1_id)  # 共享：已有的一级分类关联到该领域
    if not l2_name:
        return l1_id
    l2 = next((c for c in db.list_categories(parent_id=l1_id) if c["name"] == l2_name), None)
    return l2["id"] if l2 else db.add_category(l2_name, parent_id=l1_id)


# ---------------------------------------------------------------------- #
# 导出
# ---------------------------------------------------------------------- #
def _entry_domain_names(db, entry) -> str:
    """条目所属一级分类关联的领域名（多个用 / 分隔）"""
    root = db.category_root(entry["category_id"]) if entry["category_id"] else None
    if root is None:
        return ""
    return "/".join(d["name"] for d in db.linked_domains(root))


def export_excel(db, path, category_id=None) -> int:
    """导出全部（或指定分类子树）为 xlsx；返回导出的条目数"""
    if category_id is None:
        entries = db.list_all_entries()
    else:
        cat_ids = [c["id"] for c in _gather_categories(db, parent_id=category_id)]
        entries = [e for cid in cat_ids for e in db.list_entries(cid)]

    wb = Workbook()
    ws = wb.active
    ws.title = "提示词"
    # 2026-09-13（1-A-5 收尾 / 1-C-4）：既有 14 列不变，自定义字段列与"标签"列追加在最后
    custom_defs = [d for d in db.list_field_defs() if not d.get("is_builtin")]
    ws.append(_HEADERS + [f"{_CUSTOM_PREFIX}{d['display_name']}" for d in custom_defs]
              + [_TAG_HEADER, _GALLERY_HEADER, _PROJECT_HEADER])
    for e in entries:
        chain = _chain_names(db, e["category_id"]) if e["category_id"] else []
        ws.append([
            _entry_domain_names(db, e),
            chain[0] if len(chain) > 0 else "",
            chain[1] if len(chain) > 1 else "",
            e["name"], e["intro"], e["origin"], e["features"], e["scenes"], e["works"],
            e["image_desc"], e["prompt_cn"], e["prompt_en"], e["image_plan"], e["is_favorite"],
            *_custom_values(db, e["id"], custom_defs),
            _tags_text(db, e["id"]),
            _gallery_text(db, e["id"]),
            _entry_project_name(db, e),
        ])
    wb.save(path)
    return len(entries)


def export_payload_excel(payload, path, meta=None) -> int:
    """把**向导载荷**（本次结构化出的条目）导出为 xlsx（2026-09-14，5-d-4）。

    - 列与 `export_excel` 完全一致（含"自定义：<显示名>"、"标签"、"图集"、"项目类别"），
      因此该文件可被 `import_excel` 直接读回（可随时重导，重复导入自动去重）；
    - `meta`：与 `payload["entries"]` **一一对应**的 (项目类别, 根目录) 列表，缺省留空；
    - 返回写出的条目数。
    """
    ents = payload.get("entries") or []
    defs = [d for d in (payload.get("field_defs") or [])
            if not int(d.get("is_builtin") or 0) and d.get("field_key")]
    keys = [d["field_key"] for d in defs]
    meta = list(meta or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "提示词"
    ws.append(_HEADERS + [f"{_CUSTOM_PREFIX}{d.get('display_name') or d['field_key']}"
                          for d in defs]
              + [_TAG_HEADER, _GALLERY_HEADER, _PROJECT_HEADER])
    for i, ep in enumerate(ents):
        proj, dom = "", ""
        if i < len(meta) and meta[i]:
            pair = tuple(meta[i])
            proj = pair[0] if len(pair) > 0 else ""
            dom = pair[1] if len(pair) > 1 else ""
        path_ = [str(x) for x in (ep.get("path") or [])]
        cf = ep.get("custom_fields") or {}
        ws.append([
            dom,
            path_[0] if len(path_) > 0 else "",
            path_[1] if len(path_) > 1 else "",
            ep.get("name", ""), ep.get("intro", ""), ep.get("origin", ""),
            ep.get("features", ""), ep.get("scenes", ""), ep.get("works", ""),
            ep.get("image_desc", ""), ep.get("prompt_cn", ""), ep.get("prompt_en", ""),
            ep.get("image_plan", ""), int(ep.get("is_favorite", 0) or 0),
            *[cf.get(k, "") for k in keys],
            "、".join(ep.get("tags") or []),
            "",
            proj,
        ])
    wb.save(path)
    return len(ents)


# ---------------------------------------------------------------------- #
# 导入
# ---------------------------------------------------------------------- #
def count_excel_rows(path) -> int:
    """预览：Excel 中除表头外的数据行数（用于进度条总量）"""
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb.active
        return max((ws.max_row or 1) - 1, 0)
    finally:
        wb.close()


def import_excel(db, path, progress_cb=None) -> dict:
    """导入 xlsx；返回统计 {'entries': n}

    2026-08-29（P2 修复）：改用 read_only 流式读取，避免大文件内存暴涨。
    """
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb.active
        header = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = {h: header.index(h) for h in _HEADERS if h in header}
        if "名称" not in col:
            raise ValueError("Excel 缺少【名称】列，请使用本软件导出的模板格式")

        # 2026-09-13（1-A-5 收尾）：自定义字段列（表头 "自定义：<显示名>"）→
        # 匹配库中**同名**自定义字段；找不到同名字段则该列忽略（不影响既有 14 列导入）。
        custom_cols = []   # [(field_key, 列下标)]
        for h in header:
            if h.startswith(_CUSTOM_PREFIX):
                disp = h[len(_CUSTOM_PREFIX):].strip()
                d = next((x for x in db.list_field_defs()
                          if not x.get("is_builtin") and x["display_name"] == disp), None)
                if d:
                    custom_cols.append((d["field_key"], header.index(h)))
        # 2026-09-13（1-C-4）：标签列（表头"标签"，多个用"、"等分隔；缺失标签导入时自动创建）
        tag_col = header.index(_TAG_HEADER) if _TAG_HEADER in header else None
        # 2026-09-13（2-d）：图集列（表头"图集"；外链总是导入，本地图仅当文件存在时导入）
        gallery_col = header.index(_GALLERY_HEADER) if _GALLERY_HEADER in header else None
        # 2026-09-13（追加列收尾）：项目类别列（表头"项目类别"；缺失时按原逻辑处理）
        project_col = header.index(_PROJECT_HEADER) if _PROJECT_HEADER in header else None

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        total = max(len(rows), 1)
        done, skipped, processed = 0, 0, 0
        pending = []
        pending_cf = []   # 与 pending 一一对应的自定义字段取值（批量插入后写回）
        pending_tags = []  # 与 pending 一一对应的标签名列表（批量插入后写回）
        pending_img = []   # 2026-09-13（2-d）：与 pending 一一对应的图集（批量插入后写回）
        seen_by_cat = {}  # category_id -> 该分类现有条目的"详情内容"键集合（P1-1 去重用）
        for r in rows:
            name = str(r[col["名称"]] or "").strip()
            if not name:
                processed += 1  # 2026-09-09（P2-16）：空行也推进进度，保证进度条走满
                if progress_cb:
                    progress_cb(processed, total, "（空行，跳过）")
                continue
            cid = resolve_category_path(
                db,
                str(r[col["领域"]] or "").strip(),
                str(r[col["一级分类"]] or "").strip(),
                str(r[col["二级分类"]] or "").strip(),
                (str(r[project_col] or "").strip() if project_col is not None else ""),
            )
            e = Entry(category_id=cid, name=name,
                      intro=_cell(r, col, "介绍"), origin=_cell(r, col, "溯源"),
                      features=_cell(r, col, "核心特征"), scenes=_cell(r, col, "应用场景"),
                      works=_cell(r, col, "代表作"), image_desc=_cell(r, col, "配图描述"),
                      prompt_cn=_cell(r, col, "提示词中文"), prompt_en=_cell(r, col, "提示词英文"),
                      image_plan=_cell(r, col, "图像获取方案"),
                      is_favorite=_to_int(_cell(r, col, "收藏")))
            # 2026-08-18（P1-1）：详情内容去重——仅当内容完全相同才跳过，名称相同但内容不同仍新增
            keys = seen_by_cat.setdefault(cid, set())
            if not keys:
                keys.update(Database.content_key(x) for x in db.list_entries(cid))
            key = Database.content_key(e)
            processed += 1
            if key in keys:
                skipped += 1
                if progress_cb:
                    progress_cb(processed, total, f"{name}（重复跳过）")
                continue
            keys.add(key)
            pending.append(e)
            # 2026-09-13：收集该行的自定义字段取值（空值不写）
            _cf = {}
            for _fk, _idx in custom_cols:
                if _idx < len(r) and r[_idx] is not None and str(r[_idx]).strip():
                    _cf[_fk] = str(r[_idx]).strip()
            pending_cf.append(_cf)
            # 2026-09-13（1-C-4）：该行标签（"、"等分隔）
            if tag_col is not None and tag_col < len(r) and r[tag_col] is not None:
                pending_tags.append(_split_tags(str(r[tag_col])))
            else:
                pending_tags.append([])
            # 2026-09-13（2-d）：该行图集（"、"等分隔；外链总是导入、本地图需文件存在）
            if gallery_col is not None and gallery_col < len(r) and r[gallery_col] is not None:
                pending_img.append(_split_gallery(str(r[gallery_col])))
            else:
                pending_img.append([])
            done += 1
            if progress_cb:
                progress_cb(processed, total, name)
        # 2026-08-18（P2-4）：改为收集后批量插入，单事务提交
        if pending:
            # 2026-09-13（1-A-5 收尾）：与 JSON 导入同样按"id 递增＝插入顺序"写回自定义取值
            _max_id_before = db.conn.execute(
                "SELECT COALESCE(MAX(id), 0) FROM entries").fetchone()[0]
            db.add_entries_batch(pending)
            _restore_pending_custom_fields(db, _max_id_before, pending_cf)
            _restore_pending_tags(db, _max_id_before, pending_tags)   # 2026-09-13（1-C-4）
            _restore_pending_images(db, _max_id_before, pending_img)  # 2026-09-13（2-d）
        return {"entries": done, "skipped": skipped}
    finally:
        wb.close()


def _cell(row, col_map, key) -> str:
    """安全读取单元格文本（不存在则返回空串）"""
    idx = col_map.get(key)
    if idx is None or idx >= len(row):
        return ""
    v = row[idx]
    return "" if v is None else str(v).strip()


def _to_int(v: str) -> int:
    """收藏列容错：兼容 1 / 1.0 / '1' / 是/√ 等写法，返回 0/1"""
    if not v:
        return 0
    try:
        return 1 if float(v) else 0
    except (ValueError, TypeError):
        return 1 if v.lower() in ("是", "true", "yes", "√", "✓") else 0


# ---------------------------------------------------------------------- #
# 自测（2026-09-13）：Excel 追加"项目类别"列往返 + 旧模板兼容
# ---------------------------------------------------------------------- #
def _excel_selftest() -> None:
    """Excel 自测：导出含"项目类别"列、导入恢复项目归属、旧模板（无该列）仍可导入。"""
    import os as _os
    import shutil
    import tempfile

    from openpyxl import Workbook, load_workbook
    from ..database import Database
    from ..models import Entry

    tmp = tempfile.mkdtemp(prefix="promptsprite_excel_")
    try:
        # ---------- A. 导出：含"项目类别"列且取值正确 ---------- #
        db = Database(_os.path.join(tmp, "a.db"))
        xlsx = _os.path.join(tmp, "out.xlsx")
        try:
            pid = db.add_project("我的项目")
            did = db.add_domain("根A", project_id=pid)
            l1 = db.add_category("一级X", domain_id=did)
            l2 = db.add_category("二级Y", parent_id=l1)
            eid = db.add_entry(Entry(name="条目一", category_id=l2, prompt_cn="P1"))
            db.set_entry_tags(eid, ["标签T"])
            n = export_excel(db, xlsx)
            assert n == 1
        finally:
            db.close()
        wb = load_workbook(xlsx, read_only=True)
        try:
            ws = wb.active
            header = [str(c.value).strip() if c.value else "" for c in ws[1]]
            assert _PROJECT_HEADER in header, "导出应含『项目类别』列"
            assert header.index(_PROJECT_HEADER) == len(header) - 1, "应为最后一列（追加列）"
            row = [c.value for c in ws[2]]
            assert str(row[header.index(_PROJECT_HEADER)] or "").strip() == "我的项目"
            assert str(row[header.index("领域")] or "").strip() == "根A"
        finally:
            wb.close()
        assert count_excel_rows(xlsx) == 1

        # ---------- B. 导入到新库：项目类别归属被恢复 ---------- #
        db2 = Database(_os.path.join(tmp, "b.db"))
        try:
            res = import_excel(db2, xlsx)
            assert res["entries"] == 1 and res["skipped"] == 0
            cid = db2.find_category_by_chain(["一级X", "二级Y"])
            assert cid and [x["name"] for x in db2.list_entries(cid)] == ["条目一"]
            dom = next(d for d in db2.list_domains() if d["name"] == "根A")
            p = db2.get_project(dom["project_id"]) if dom["project_id"] else None
            assert p and p["name"] == "我的项目", "应恢复『项目类别』归属"
        finally:
            db2.close()

        # ---------- C. 旧模板（无"项目类别"列）仍可导入 ---------- #
        legacy = _os.path.join(tmp, "legacy.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.append(_HEADERS)
        ws.append(["老根", "老一级", "", "老条目", "", "", "", "", "", "", "P", "", "", 0])
        wb.save(legacy)
        db3 = Database(_os.path.join(tmp, "c.db"))
        try:
            res3 = import_excel(db3, legacy)
            assert res3["entries"] == 1
            cid = db3.find_category_by_chain(["老一级"])
            assert cid and [x["name"] for x in db3.list_entries(cid)] == ["老条目"]
            dom = next(d for d in db3.list_domains() if d["name"] == "老根")
            p = db3.get_project(dom["project_id"]) if dom["project_id"] else None
            assert p and p["name"] == "未明确分类", "旧模板新建根目录应落兜底项目类别"
        finally:
            db3.close()

        # ---------- D. 重复导入判重 ---------- #
        db4 = Database(_os.path.join(tmp, "d.db"))
        try:
            import_excel(db4, xlsx)
            res4 = import_excel(db4, xlsx)
            assert res4["entries"] == 0 and res4["skipped"] == 1
        finally:
            db4.close()

        print("[Excel] 追加『项目类别』列（导出取值正确·导入恢复归属）/旧模板兼容/"
              "判重幂等 通过")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    _excel_selftest()
